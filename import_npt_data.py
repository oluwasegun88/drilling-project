import sys
from pathlib import Path
from datetime import datetime, date

# Path bootstrap
workspace_root = Path(__file__).parent
sys.path.insert(0, str(workspace_root))

import openpyxl
from firebase_database.setup_firebase import db

# Configuration
EXCEL_FILE = Path(__file__).parent / "Awoba NW 5 NPT_Drilling_Performance_Dashboard.xlsx"
WELL_NAME  = "Awoba NW 5"

# Lookup values
PHASES = [
    "Conductor Piling/Cleanout",
    "17-1/2 x 23 Hole Section",
    "16 Hole Section",
    "12 1/4 Hole Section",
    "Perforation and Wellbore Cleanout",
    "Completions",
]

CATEGORIES = [
    "Mechanical", "Electrical", "Wellbore", "BHA / Tools",
    "Mud / Solids", "Cementing", "Logistics", "Weather",
    "HSE", "Waiting on Service", "Human Error",
    "Equipment Failure", "Wellhead", "Other",
]

RESPONSIBLE_PARTIES = [
    "SMS Joy", "Newcross EP", "OMASUP", "MultiChase",
    "Clinton", "Marine/Logistics", "TBD", "Fredrikov",
    "Mega Field", "Lagrange", "Ulla", "Geowell",
]

WELLS = ["Awoba NW 5"]


def date_to_str(val) -> str:
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    elif isinstance(val, (int, float)):
        d = date.fromtimestamp((val - 25569) * 86400)
        return d.strftime("%Y-%m-%d")
    return str(val)


def safe_float(value) -> float:
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def seed_collection(collection_name, items, label):
    print(f"\nSeeding {label}...")
    collection = db.collection(collection_name)
    existing = set()
    for doc in collection.stream():
        data = doc.to_dict()
        if data.get("name"):
            existing.add(data["name"].strip().lower())

    added = skipped = 0
    for item in items:
        if item.strip().lower() in existing:
            skipped += 1
        else:
            db.collection(collection_name).document().set({"name": item, "description": ""})
            print(f"   Added: {item}")
            added += 1
    print(f"   Done - {added} added, {skipped} skipped")


def import_npt_records():
    print(f"\nReading Excel file: {EXCEL_FILE}")

    if not EXCEL_FILE.exists():
        print(f"ERROR: Excel file not found at: {EXCEL_FILE}")
        return

    wb = openpyxl.load_workbook(str(EXCEL_FILE), read_only=True, data_only=True)

    if "NPT" not in wb.sheetnames:
        print(f"ERROR: NPT sheet not found. Sheets: {wb.sheetnames}")
        return

    ws         = wb["NPT"]
    collection = db.collection("nptRecords")

    # Get existing records to avoid duplicates
    existing_keys = set()
    for doc in collection.stream():
        data = doc.to_dict()
        key  = f"{data.get('date')}_{data.get('phase')}_{data.get('hours')}"
        existing_keys.add(key)

    added = skipped = errors = 0

    print(f"\nImporting NPT records for well: {WELL_NAME}")

    for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
        # Row 1 = title, Row 2 = empty, Row 3 = headers, Row 4+ = data
        if row_num <= 3:
            continue

        # Data starts at column index 1 (column B in Excel)
        # Row structure: [None, date, phase, hours, days, npt_cost, description, category, responsible_party]
        if not row[1]:  # date is in column index 1
            continue

        try:
            date_val    = row[1]
            phase       = str(row[2]).strip() if row[2] else ""
            hours       = safe_float(row[3])
            npt_cost    = safe_float(row[5])
            description = str(row[6]).strip() if row[6] else ""
            category    = str(row[7]).strip() if row[7] else ""
            responsible = str(row[8]).strip() if row[8] else ""

            if not phase or not category:
                continue

            date_str = date_to_str(date_val)
            key      = f"{date_str}_{phase}_{hours}"

            if key in existing_keys:
                skipped += 1
                continue

            record = {
                "well":              WELL_NAME,
                "date":              date_str,
                "phase":             phase,
                "hours":             hours,
                "days":              round(hours / 24, 6),
                "npt_cost":          round(npt_cost, 2),
                "description":       description,
                "category":          category,
                "responsible_party": responsible,
            }

            collection.document().set(record)
            existing_keys.add(key)
            added += 1
            print(f"   Row {row_num} | {date_str} | {phase[:25]} | {hours}hrs | {category}")

        except Exception as e:
            errors += 1
            print(f"   Row {row_num} error: {e}")

    wb.close()
    print(f"\nNPT Import complete:")
    print(f"   Added:   {added} records")
    print(f"   Skipped: {skipped} records")
    print(f"   Errors:  {errors} records")


if __name__ == "__main__":
    print("=" * 60)
    print("  RADA Analysis - NPT Data Import Script")
    print("=" * 60)

    seed_collection("wells",              WELLS,               "Wells")
    seed_collection("phases",             PHASES,              "Phases")
    seed_collection("categories",         CATEGORIES,          "Categories")
    seed_collection("responsibleParties", RESPONSIBLE_PARTIES, "Responsible Parties")

    import_npt_records()

    print("\n" + "=" * 60)
    print("  Import complete! Test your API:")
    print("  GET http://127.0.0.1:8000/drilling/npt/")
    print("  GET http://127.0.0.1:8000/drilling/dashboard/metrics?well=Awoba NW 5")
    print("=" * 60)